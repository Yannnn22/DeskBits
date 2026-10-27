"""文件工具层：file_diff / file_patch + /api/files 浏览 API。"""

import asyncio
from pathlib import Path

from fastapi.testclient import TestClient

from desktopmate_agent.agent.intent import HeuristicIntentRouter, IntentType
from desktopmate_agent.config import Settings
from desktopmate_agent.main import create_app
from desktopmate_agent.tools.file_tools import FileDiffTool, FilePatchTool

AUTH_HEADERS = {"X-DesktopMate-Token": "test-token"}


def make_settings(tmp_path: Path, root: Path) -> Settings:
    return Settings(
        trace_dir=tmp_path / "traces",
        settings_dir=tmp_path / "settings",
        app_data_dir=tmp_path,
        knowledge_dir=tmp_path / "knowledge",
        knowledge_db_path=tmp_path / "desktopmate.db",
        session_db_path=tmp_path / "desktopmate.db",
        auth_token="test-token",
        openai_api_key=None,
        deepseek_api_key=None,
        embedding_provider="fake",
        embedding_dimensions=16,
        allowed_roots=str(root),
    )


# --- file_diff ---


def test_file_diff_preview(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    target = root / "note.md"
    target.write_text("hello\nworld\n", encoding="utf-8")
    tool = FileDiffTool(allowed_roots=[root])

    result = asyncio.run(tool.run(FileDiffTool.args_schema(path=str(target), proposed_content="hello\nuniverse\n")))
    assert result.status == "success"
    assert result.data["changed"] is True
    assert "-world" in result.data["diff"] and "+universe" in result.data["diff"]
    assert target.read_text(encoding="utf-8") == "hello\nworld\n"  # 未写入


def test_file_diff_no_change(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    target = root / "note.md"
    target.write_text("same\n", encoding="utf-8")
    tool = FileDiffTool(allowed_roots=[root])
    result = asyncio.run(tool.run(FileDiffTool.args_schema(path=str(target), proposed_content="same\n")))
    assert result.status == "success"
    assert result.data["changed"] is False
    assert result.data["diff"] == ""


def test_file_diff_out_of_root(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    outside = tmp_path / "secret.txt"
    outside.write_text("x", encoding="utf-8")
    tool = FileDiffTool(allowed_roots=[root])
    result = asyncio.run(tool.run(FileDiffTool.args_schema(path=str(outside), proposed_content="y")))
    assert result.status == "error"


# --- file_patch ---


def test_file_patch_create_and_update(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    tool = FilePatchTool(allowed_roots=[root])
    target = root / "new.md"

    created = asyncio.run(tool.run(FilePatchTool.args_schema(path=str(target), new_content="line1\n")))
    assert created.status == "success"
    assert created.data["created"] is True
    assert target.read_text(encoding="utf-8") == "line1\n"

    updated = asyncio.run(tool.run(FilePatchTool.args_schema(path=str(target), new_content="line1\nline2\n")))
    assert updated.status == "success"
    assert updated.data["created"] is False
    assert "+line2" in updated.data["diff"]
    assert target.read_text(encoding="utf-8") == "line1\nline2\n"
    # 无临时文件残留
    assert list(root.glob(".*.tmp")) == []


def test_file_patch_out_of_root(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    outside = tmp_path / "secret.txt"
    tool = FilePatchTool(allowed_roots=[root])
    result = asyncio.run(tool.run(FilePatchTool.args_schema(path=str(outside), new_content="x")))
    assert result.status == "error"
    assert not outside.exists()


def test_file_patch_is_permission_gated():
    tool = FilePatchTool(allowed_roots=[Path("/tmp")])
    assert tool.requires_permission is True
    assert tool.risk_level.value == "write"


def test_file_patch_approval_preview_shows_diff_without_writing(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    target = root / "note.md"
    target.write_text("before\n", encoding="utf-8")
    tool = FilePatchTool(allowed_roots=[root])

    preview = asyncio.run(
        tool.approval_preview({"path": str(target), "new_content": "after\n"})
    )

    assert preview == {
        "kind": "file_patch",
        "ready": True,
        "path": str(target),
        "created": False,
        "changed": True,
        "diff": f"--- {target}\n+++ {target}\n@@ -1 +1 @@\n-before\n+after\n",
    }
    assert target.read_text(encoding="utf-8") == "before\n"


# --- 意图路由 ---


def test_intent_diff_and_patch():
    router = HeuristicIntentRouter()

    d = router.route("/diff /tmp/a.md\nnew content")
    assert d.type == IntentType.TOOL_CALL
    assert d.tool_name == "file_diff"
    assert d.tool_args["path"] == "/tmp/a.md"
    assert d.tool_args["proposed_content"] == "new content"

    p = router.route("/patch /tmp/a.md\nupdated")
    assert p.type == IntentType.TOOL_CALL
    assert p.tool_name == "file_patch"
    assert p.tool_args["new_content"] == "updated"

    bad = router.route("/diff ")
    assert bad.type == IntentType.REFUSAL

    read = router.route("/read /tmp/a.md")
    assert read.tool_name == "read_file"


# --- /api/files 浏览 API ---


def test_files_preview():
    """/api/files/preview：文本/docx/pptx/pdf 提取 + 不支持格式 + 越权。"""
    import tempfile
    import zipfile

    root = Path(tempfile.mkdtemp(prefix="preview-test-"))
    settings = make_settings(root, root)

    # 文本
    md = root / "note.md"
    md.write_text("# 标题\n内容行", encoding="utf-8")

    # docx：手工构造最小 docx（document.xml）
    def minimal_docx(path: Path) -> None:
        body = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body>"
            '<w:p><w:r><w:t>第一段文档内容</w:t></w:r></w:p>'
            "</w:body></w:document>"
        ).encode("utf-8")
        doc_zip = {
            "[Content_Types].xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                "</Types>"
            ).encode("utf-8"),
            "_rels/.rels": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                "</Relationships>"
            ).encode("utf-8"),
            "word/document.xml": body,
        }
        with zipfile.ZipFile(path, "w") as zf:
            for name, data in doc_zip.items():
                zf.writestr(name, data)

    docx_path = root / "doc.docx"
    minimal_docx(docx_path)

    # pptx：用 python-pptx 生成
    from pptx import Presentation
    from pptx.util import Emu

    ppt_path = root / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    slide.shapes.add_textbox(Emu(0), Emu(0), Emu(4000000), Emu(2000000)).text_frame.text = "第一页要点"
    prs.save(str(ppt_path))

    # 不支持格式（旧版 .doc 魔数）
    legacy = root / "old.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 legacy")

    with TestClient(create_app(settings=settings)) as client:
        resp = client.get(f"/api/files/preview?path={md}", headers=AUTH_HEADERS)
        assert resp.status_code == 200
        data = resp.json()
        assert data["kind"] == "text" and data["name"] == "note.md" and "标题" in data["content"]

        pdf_docx = client.get(f"/api/files/preview?path={docx_path}", headers=AUTH_HEADERS)
        assert pdf_docx.status_code == 200
        assert pdf_docx.json()["kind"] == "docx" and "第一段文档内容" in pdf_docx.json()["content"]

        ppt = client.get(f"/api/files/preview?path={ppt_path}", headers=AUTH_HEADERS)
        assert ppt.status_code == 200
        assert ppt.json()["kind"] == "pptx" and "第一页要点" in ppt.json()["content"]

        bad = client.get(f"/api/files/preview?path={legacy}", headers=AUTH_HEADERS)
        assert bad.status_code == 415
        assert "docx/pptx" in bad.json()["error"]["message"]

        missing = client.get("/api/files/preview?path=/nonexistent/x.md", headers=AUTH_HEADERS)
        assert missing.status_code == 403  # 不在允许根目录内


def test_files_api_list_and_read(tmp_path):
    root = tmp_path / "root"
    (root / "sub").mkdir(parents=True)
    (root / "a.txt").write_text("alpha\n", encoding="utf-8")
    (root / "b.md").write_text("# beta\n", encoding="utf-8")
    settings = make_settings(tmp_path, root)
    client = TestClient(create_app(settings=settings))

    listing = client.get("/api/files/list", params={"path": str(root)}, headers=AUTH_HEADERS)
    assert listing.status_code == 200
    body = listing.json()
    assert body["path"] == str(root.resolve())
    names = [e["name"] for e in body["entries"]]
    assert names == ["sub", "a.txt", "b.md"]  # 目录在前，文件按名称排序
    assert body["entries"][1]["type"] == "file"

    reading = client.get("/api/files/read", params={"path": str(root / "a.txt")}, headers=AUTH_HEADERS)
    assert reading.status_code == 200
    assert reading.json()["content"] == "alpha\n"

    missing = client.get("/api/files/read", params={"path": str(root / "nope.txt")}, headers=AUTH_HEADERS)
    assert missing.status_code == 404

    traversal = client.get("/api/files/list", params={"path": str(tmp_path)}, headers=AUTH_HEADERS)
    assert traversal.status_code == 403  # root 之外被拒

    no_auth = client.get("/api/files/list", params={"path": str(root)})
    assert no_auth.status_code == 401 or no_auth.status_code == 403


def test_files_tools_registered_in_app(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    settings = make_settings(tmp_path, root)
    client = TestClient(create_app(settings=settings))
    tools = client.get("/api/tools", headers=AUTH_HEADERS).json()["tools"]
    names = {t.get("name") for t in tools}
    assert {"file_diff", "file_patch", "read_file"} <= names


# --- F1: SearchFilesTool 无权限目录不应崩溃（B2 修复） ---


def test_search_files_skips_unreadable_dirs(tmp_path):
    """遍历到无权限目录时跳过而非抛 PermissionError 崩溃。"""
    import os

    from desktopmate_agent.tools.file_tools import SearchFilesTool

    root = tmp_path / "root"
    (root / "visible").mkdir(parents=True)
    (root / "visible" / "target.txt").write_text("x", encoding="utf-8")

    locked = root / "locked"
    locked.mkdir()
    (locked / "secret.txt").write_text("y", encoding="utf-8")
    os.chmod(locked, 0o000)  # 无权限目录（POSIX）
    try:
        tool = SearchFilesTool(allowed_roots=[root])
        result = asyncio.run(tool.run(SearchFilesTool.args_schema(query="txt", limit=10)))
        assert result.status == "success"  # 不崩溃
        assert any("target.txt" in m for m in result.data["matches"])
        assert all("secret.txt" not in m for m in result.data["matches"])
    finally:
        os.chmod(locked, 0o755)  # 恢复权限，避免 tmp 清理失败
