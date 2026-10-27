from hashlib import sha1
from pathlib import Path

from desktopmate_agent.rag.models import KnowledgeDocument


class TextDocumentLoader:
    supported_extensions = {".md", ".markdown", ".txt"}

    def load(self, path: Path) -> KnowledgeDocument:
        if path.suffix.lower() not in self.supported_extensions:
            raise ValueError(f"unsupported document type: {path.suffix}")
        return self.load_text(
            text=path.read_text(encoding="utf-8"),
            source_path=str(path),
        )

    def load_text(self, *, text: str, source_path: str) -> KnowledgeDocument:
        title = Path(source_path).stem or "untitled"
        document_id = sha1(f"{source_path}\n{text}".encode("utf-8")).hexdigest()
        return KnowledgeDocument(
            document_id=document_id,
            source_path=source_path,
            title=title,
            text=text,
            metadata={"extension": Path(source_path).suffix.lower()},
        )


class PdfDocumentLoader:
    supported_extensions = {".pdf"}

    def load(self, path: Path) -> KnowledgeDocument:
        if path.suffix.lower() not in self.supported_extensions:
            raise ValueError(f"unsupported document type: {path.suffix}")
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(page for page in pages if page.strip())
        if not text.strip():
            raise ValueError(f"PDF 无可提取文本: {path.name}")
        return self._from_text(text, str(path))

    def _from_text(self, text: str, source_path: str) -> KnowledgeDocument:
        title = Path(source_path).stem or "untitled"
        document_id = sha1(f"{source_path}\n{text}".encode("utf-8")).hexdigest()
        return KnowledgeDocument(
            document_id=document_id,
            source_path=source_path,
            title=title,
            text=text,
            metadata={"extension": ".pdf"},
        )


class DocxDocumentLoader:
    supported_extensions = {".docx"}

    def load(self, path: Path) -> KnowledgeDocument:
        if path.suffix.lower() not in self.supported_extensions:
            raise ValueError(f"unsupported document type: {path.suffix}")
        from docx import Document

        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        text = "\n".join(paragraphs)
        if not text.strip():
            raise ValueError(f"DOCX 无可提取文本: {path.name}")
        return self._from_text(text, str(path))

    def _from_text(self, text: str, source_path: str) -> KnowledgeDocument:
        title = Path(source_path).stem or "untitled"
        document_id = sha1(f"{source_path}\n{text}".encode("utf-8")).hexdigest()
        return KnowledgeDocument(
            document_id=document_id,
            source_path=source_path,
            title=title,
            text=text,
            metadata={"extension": ".docx"},
        )
